"""Cotações da Scot Consultoria (boi, vaca, novilha, reposição, grãos, leite e outras).

Lê as páginas de https://www.scotconsultoria.com.br/cotacoes/ (HTML simples, sem
login), transforma cada tabela em linhas no formato longo (uma linha por data,
praça e indicador) e acumula o histórico, porque o site só mostra a cotação mais
recente.

Uso:
    python scot_cotacoes.py                  # coleta todas as páginas
    python scot_cotacoes.py --paginas boi-gordo reposicao
    python scot_cotacoes.py --html testes/html --sem-gravar   # testa com páginas salvas
    python scot_cotacoes.py --sem-excel      # só o CSV (usado no GitHub Actions)

Arquivos gerados (na pasta deste script):
    scot_cotacoes.xlsx           planilha com uma aba por grupo (refeita a cada execução)
    dados/scot_cotacoes_AAAA.csv histórico no formato longo, um arquivo por ano (lido pelo Power BI)
    dados/diario/AAAA-MM-DD.csv  linhas de cada data de cotação (fonte das saídas)
    dados/controle.json          quando cada grupo de cada data foi coletado e revisado

Regras:
    - A data de cada linha é a que a Scot publica no título da tabela, não a data da
      coleta. A Scot publica a cotação de um dia de madrugada (entre ~22h30 e ~6h25);
      durante o dia o site mostra a do dia útil anterior.
    - Uma data já gravada só é regravada se os valores mudarem (a Scot às vezes
      corrige a tabela depois de publicada); o controle registra cada revisão.
    - Se a estrutura de uma tabela mudar (títulos, colunas ou poucas linhas), a
      página dá erro e nada dela é gravado; as demais páginas seguem.
    - Código de saída: 0 tudo certo; 1 alguma página falhou nesta execução (ex.: o
      site não respondeu; é tentada de novo na próxima); 2 algum grupo está há dias
      sem data nova (ver ATRASO_MAXIMO_DIAS), o que pede atenção.
    - No Mercado futuro, a data é a do último ajuste da tabela (a Scot publica no
      dia D os ajustes da B3 de D-2 e D-1). No Leite, a data é o 1º dia do mês do
      pagamento (a cotação é mensal).

Requer Python 3.10+ com lxml e openpyxl.
"""

from __future__ import annotations

import argparse
import collections
import csv
import datetime as dt
import http.client
import json
import re
import sys
import time
import unicodedata
import urllib.error
import urllib.request
from pathlib import Path

import lxml.html
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = Path(__file__).resolve().parent
DIR_DIARIO = BASE / 'dados' / 'diario'
ARQ_CONTROLE = BASE / 'dados' / 'controle.json'
ARQ_EXCEL = BASE / 'scot_cotacoes.xlsx'
ARQ_CSV = BASE / 'dados' / 'scot_cotacoes.csv'

# A Scot publica no horário de Brasília (sem horário de verão desde 2019);
# servidores como o do GitHub Actions rodam em UTC.
FUSO_BRASILIA = dt.timezone(dt.timedelta(hours=-3))


def agora() -> dt.datetime:
    return dt.datetime.now(FUSO_BRASILIA).replace(microsecond=0)


def hoje() -> dt.date:
    return agora().date()

URL_PAGINA = 'https://www.scotconsultoria.com.br/cotacoes/{}/'
USER_AGENT = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
              '(KHTML, like Gecko) Chrome/125.0.0.0 Safari/537.36')
TENTATIVAS = 4
ESPERA_INICIAL_S = 10
PAUSA_ENTRE_PAGINAS_S = 3


class Bloqueado(Exception):
    """O site recusou o acesso (HTTP 403 ou conexão derrubada): o firewall da Scot às vezes
    bloqueia o endereço de um servidor do GitHub. Insistir piora; a próxima coleta (outro
    servidor) tenta de novo."""
# Dias corridos sem data nova a partir dos quais o grupo é dado como atrasado (o
# site caiu de vez ou mudou o layout). Folga para fins de semana e feriados.
ATRASO_MAXIMO_DIAS = 5
# No Leite a data é o 1º dia do mês do pagamento, e a tabela só é trocada no mês seguinte
# ao do pagamento (a de setembro/26 ainda estava no site em 23/10/2026, com 52 dias).
ATRASO_MAXIMO_DIAS_GRUPO = {'Couro e sebo': 10, 'Leite': 100}

# Ordem dos grupos nas saídas e nas abas da planilha.
GRUPOS = ['Boi gordo', 'Boi China', 'Vaca gorda', 'Novilha gorda', 'Indicador do boi gordo',
          'Reposição', 'Mercado futuro do boi', 'Milho', 'Soja', 'Atacado', 'Equivalentes',
          'Boi no mundo', 'Couro e sebo', 'Leite']

# Notas da Scot (rodapés das tabelas), repetidas na aba "Sobre" da planilha.
NOTAS = {
    'Boi gordo': 'Preços brutos e descontados, à vista e a prazo (30 dias). "desc. Funrural": com o desconto '
                 'do Funrural; "desc. Senar": sem o desconto do Funrural, mas com o desconto de 0,2% do Senar. '
                 'Diferencial de base: em relação a São Paulo (%). No RS o preço é por kg vivo (50% de rendimento '
                 'de carcaça). MT Cuiabá inclui Rondonópolis.',
    'Boi China': 'Boi China a prazo (30 dias): bovino com no máximo quatro dentes incisivos permanentes e idade '
                 'inferior a 30 meses no abate, com GTA especificando a idade compatível. Líquido: com o desconto '
                 'do Funrural.',
    'Vaca gorda': 'Mesmas colunas do boi gordo, sem o diferencial de base.',
    'Novilha gorda': 'Mesmas colunas do boi gordo, sem o diferencial de base.',
    'Indicador do boi gordo': 'Média móvel dos últimos cinco dias úteis de todas as cotações do boi gordo '
                              'levantadas pela Scot, praça por praça. Variação = indicador do dia - do dia anterior.',
    'Reposição': 'R$ por cabeça, R$ por kg e relação de troca (arrobas de boi gordo necessárias para comprar '
                 'um animal). Na tabela de nelore, SC e RS são bovinos cruzados com o peso de referência do nelore.',
    'Mercado futuro do boi': 'Ajustes do boi gordo na B3 (R$/@ à vista). A data é a do último ajuste. Câmbio e '
                             'US$ à vista: projeção do físico no futuro. "Contratos negociados" é a coluna que a '
                             'Scot chama de "C. Abertos" (o número confere com os negociados no boletim da B3).',
    'Milho': 'Preço médio de compra, R$/saca de 60 kg, sem frete, à vista, para entrega imediata (fonte AgRural).',
    'Soja': 'Preço médio de compra, R$/saca de 60 kg, sem frete, à vista, para entrega imediata (fonte AgRural).',
    'Atacado': 'Atacado de São Paulo, R$/kg, à vista, sem bonificação e sem ICMS.',
    'Equivalentes': 'Equivalente físico: 48% traseiro + 39% dianteiro + 13% ponta de agulha. Equivalente Scot: '
                    'físico + couro + sebo. Carcaça/desossa: carne com/sem osso + couro + sebo + miúdos + derivados '
                    '+ subprodutos. Margem de comercialização em relação à arroba recebida pelo pecuarista (não é '
                    'lucratividade).',
    'Boi no mundo': 'Preço da arroba do boi gordo em outros países, US$/@.',
    'Couro e sebo': 'R$/kg, FOB, sem bonificação e sem ICMS. Sebo a prazo; couro verde à vista. Levantamento '
                    'semanal, atualizado às quintas-feiras.',
    'Leite': 'Preço do leite pago ao produtor, sem frete, mínimo e máximo por qualidade (com bonificação por '
             'qualidade) e padrão, e médias regionais. Mensal: a data é o mês do pagamento (produção do mês '
             'anterior). Valores 0,000 na tabela da Scot (sem levantamento) ficam fora.',
}

UFS = set('AC AL AP AM BA CE DF ES GO MA MT MS MG PA PB PR PE PI RJ RN RS RO RR SC SP SE TO'.split())
ESTADOS = {
    'Acre': 'AC', 'Alagoas': 'AL', 'Amapá': 'AP', 'Amazonas': 'AM', 'Bahia': 'BA', 'Ceará': 'CE',
    'Distrito Federal': 'DF', 'Espírito Santo': 'ES', 'Goiás': 'GO', 'Maranhão': 'MA',
    'Mato Grosso do Sul': 'MS', 'Mato Grosso': 'MT', 'Minas Gerais': 'MG', 'Pará': 'PA',
    'Paraíba': 'PB', 'Paraná': 'PR', 'Pernambuco': 'PE', 'Piauí': 'PI', 'Rio de Janeiro': 'RJ',
    'Rio Grande do Norte': 'RN', 'Rio Grande do Sul': 'RS', 'Rondônia': 'RO', 'Roraima': 'RR',
    'Santa Catarina': 'SC', 'São Paulo': 'SP', 'Sergipe': 'SE', 'Tocantins': 'TO',
}
MESES_ABREV = dict(zip('jan fev mar abr mai jun jul ago set out nov dez'.split(), range(1, 13)))
MESES = dict(zip(('janeiro fevereiro março abril maio junho julho agosto setembro outubro '
                  'novembro dezembro').split(), range(1, 13)))
# Setas das tabelas: n = estável, s = subiu, d e t = desceu (t só no indicador).
TENDENCIAS = {'n': 'Estável', 's': 'Subiu', 'd': 'Desceu', 't': 'Desceu'}

CAMPOS_CSV = ['data', 'grupo', 'categoria', 'produto', 'uf', 'praca', 'indicador',
              'unidade', 'valor', 'tendencia', 'coletado_em', 'origem']
COLUNAS_CSV = ['Data', 'Grupo', 'Categoria', 'Produto', 'UF', 'Praça', 'Indicador',
               'Unidade', 'Valor', 'Tendência', 'Coletado em', 'Origem']
ORIGEM_PROPRIA = 'Coleta própria'
DESCRICAO_COLUNAS = [
    'Data da cotação publicada pela Scot (no Leite, 1º dia do mês do pagamento)',
    'Tabela da Scot de onde a linha veio',
    'Subdivisão do grupo (ex.: Macho nelore na reposição, vencimento no mercado futuro)',
    'Animal, corte ou produto cotado',
    'Sigla do estado, quando a praça é de um estado',
    'Praça, cidade, região ou país, como a Scot escreve (sem os asteriscos de nota)',
    'O que o valor mede (ex.: À vista (bruto), R$/cab, Compra)',
    'Unidade do valor',
    'Valor numérico (ponto decimal no CSV)',
    'Seta da Scot na tabela: Estável, Subiu ou Desceu (só onde o site mostra)',
    'Data e hora (Brasília) da coleta que gravou a linha',
    'Coleta própria (lida do site por este repositório) ou o projeto de onde veio o histórico importado '
    '(Gado-Scraper, boiapi: dez/2025 a out/2026, com a data corrigida; ver historico/README.md)',
]


def log(msg: str) -> None:
    print(f'[{agora():%H:%M:%S}] {msg}', flush=True)


# --------------------------------------------------------------------------- download

def baixar(pagina: str, timeout: int = 30) -> str:
    req = urllib.request.Request(URL_PAGINA.format(pagina), headers={
        'User-Agent': USER_AGENT,
        'Accept': 'text/html,application/xhtml+xml',
        'Accept-Language': 'pt-BR,pt;q=0.9',
    })
    for tentativa in range(1, TENTATIVAS + 1):
        try:
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                conteudo = resp.read()
            try:
                return conteudo.decode('utf-8')
            except UnicodeDecodeError:
                return conteudo.decode('latin-1')
        except urllib.error.HTTPError as e:
            if e.code == 403:
                raise Bloqueado(f'HTTP 403') from e
            if e.code not in (429, 500, 502, 503, 504) or tentativa == TENTATIVAS:
                raise
            erro = f'HTTP {e.code}'
        except (urllib.error.URLError, TimeoutError, ConnectionError) as e:
            motivo = getattr(e, 'reason', e)
            if isinstance(e, (ConnectionResetError, http.client.RemoteDisconnected)) or isinstance(
                    motivo, (ConnectionResetError, http.client.RemoteDisconnected)):
                raise Bloqueado(f'conexão derrubada ({motivo})') from e
            if tentativa == TENTATIVAS:
                raise
            erro = str(motivo)
        espera = ESPERA_INICIAL_S * 3 ** (tentativa - 1)
        log(f'  {pagina}: falha ao baixar ({erro}); nova tentativa em {espera}s')
        time.sleep(espera)
    raise RuntimeError('inalcançável')


# --------------------------------------------------------------------------- leitura do HTML

def _texto(el) -> str:
    return ' '.join(el.text_content().split())


def _celulas(tr) -> list:
    return [c for c in tr if c.tag in ('td', 'th')]


def _linhas(tabela, classe: str | None = '') -> list:
    """Linhas da própria tabela (sem as de tabelas aninhadas); filtra pela classe se pedida."""
    trs = tabela.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr')
    return trs if classe == '' else [tr for tr in trs if tr.get('class') == classe]


def _num(texto: str) -> float | None:
    texto = texto.replace('%', '').replace('*', '').strip()
    if texto in ('', '-', '–'):
        return None
    return float(texto.replace('.', '').replace(',', '.'))


def _data(texto: str) -> str:
    dia, mes, ano = re.search(r'(\d{2})/(\d{2})/(\d{4})', texto).groups()
    return dt.date(int(ano), int(mes), int(dia)).isoformat()


def _tendencia(tr) -> str:
    for img in tr.iter('img'):
        m = re.search(r'/([a-z])\.gif$', img.get('src', ''))
        if m and m.group(1) in TENDENCIAS:
            return TENDENCIAS[m.group(1)]
    return ''


def _uf(praca: str) -> str:
    sigla = praca.split(' ')[0]
    if sigla in UFS:
        return sigla
    for nome in sorted(ESTADOS, key=len, reverse=True):
        if praca.startswith(nome):
            return ESTADOS[nome]
    return ''


def _limpar(texto: str) -> str:
    return texto.replace('*', '').strip()


def _tabela(doc, padrao: str):
    """Primeira tabela cujo primeiro título/célula casa com o padrão; devolve (tabela, match)."""
    regex = re.compile(padrao, re.IGNORECASE)
    for tabela in doc.iter('table'):
        celulas = tabela.xpath('.//th|.//td')
        if celulas:
            m = regex.search(_texto(celulas[0]))
            if m:
                return tabela, m
    raise ValueError(f'tabela "{padrao}" não encontrada (o layout do site mudou?)')


def _exigir(condicao: bool, mensagem: str) -> None:
    if not condicao:
        raise ValueError(mensagem)


def _registro(data, grupo, indicador, valor, unidade, *, categoria='', produto='',
              uf='', praca='', tendencia='') -> dict:
    return {'data': data, 'grupo': grupo, 'categoria': categoria, 'produto': produto,
            'uf': uf, 'praca': praca, 'indicador': indicador, 'unidade': unidade,
            'valor': valor, 'tendencia': tendencia}


# --------------------------------------------------------------------------- páginas

def _mercado_fisico(doc, grupo: str, rotulo: str, minimo: int = 25) -> list[dict]:
    """Boi gordo, vaca gorda e novilha gorda: brutos, descontados e (no boi) a base."""
    tabela, m = _tabela(doc, r'^Mercado F[íi]sico - (\d{2}/\d{2}/\d{4})$')
    data = _data(m.group(1))
    linhas = _linhas(tabela)
    subtitulos = [tr for tr in linhas if tr.get('class') == 'subtitle']
    _exigir(subtitulos and _texto(_celulas(subtitulos[0])[0]).upper() == rotulo,
            f'tabela do mercado físico não é de {rotulo}')
    cabecalho = next((tr for tr in linhas if tr.get('class') is None
                      and any(_texto(c).lower() == 'à vista' for c in _celulas(tr))), None)
    _exigir(cabecalho is not None, 'cabeçalho "à vista / 30 D" não encontrado')

    # Os pares "à vista / 30 D" vêm na ordem: bruto, desc. Funrural, desc. Senar.
    tipos = iter(['bruto', 'desc. Funrural', 'desc. Senar'])
    colunas, tipo = [], None
    for c in _celulas(cabecalho):
        nome = _texto(c).lower()
        if nome == 'à vista':
            tipo = next(tipos)
            colunas.append(('À vista', tipo))
        elif nome == '30 d':
            colunas.append(('30 dias', tipo))
        elif nome.startswith('# base'):
            colunas.append(('base', None))
        elif nome == '':
            colunas.append(None)
        else:
            raise ValueError(f'coluna inesperada no mercado físico: "{nome}"')
    _exigir(tipo == 'desc. Senar', 'colunas do mercado físico mudaram')

    registros = []
    for tr in _linhas(tabela, 'conteudo'):
        celulas = _celulas(tr)
        _exigir(len(celulas) == len(colunas) + 1,
                f'{_texto(celulas[0])}: {len(celulas) - 1} colunas, esperado {len(colunas)}')
        original = _texto(celulas[0])
        praca = _limpar(original.replace('(kg)', ''))
        unidade = 'R$/kg vivo' if '(kg)' in original else 'R$/@'
        tendencia = _tendencia(tr)
        for coluna, celula in zip(colunas, celulas[1:]):
            if coluna is None:
                continue
            valor = _num(_texto(celula))
            if valor is None:
                continue
            if coluna[0] == 'base':
                registros.append(_registro(data, grupo, 'Diferencial de base SP', valor, '%',
                                           produto=grupo, uf=_uf(praca), praca=praca))
            else:
                registros.append(_registro(data, grupo, f'{coluna[0]} ({coluna[1]})', valor, unidade,
                                           produto=grupo, uf=_uf(praca), praca=praca,
                                           tendencia=tendencia))
    _exigir(len({r['praca'] for r in registros}) >= minimo, f'{grupo}: poucas praças')
    return registros


def _boi_china(doc) -> list[dict]:
    tabela, m = _tabela(doc, r'^Boi China a Prazo \(R\$/@\) - (\d{2}/\d{2}/\d{4})$')
    data = _data(m.group(1))
    registros = []
    for tr in _linhas(tabela, 'conteudo'):
        celulas = [_texto(c) for c in _celulas(tr)]
        if celulas[0] == 'UF':
            _exigir('bruto' in celulas[1].lower() and 'quido' in celulas[2].lower(),
                    'colunas do Boi China mudaram')
            continue
        praca = _limpar(celulas[0])
        for indicador, texto in zip(['Bruto 30 dias', 'Líquido 30 dias'], celulas[1:3]):
            valor = _num(texto)
            if valor is not None:
                registros.append(_registro(data, 'Boi China', indicador, valor, 'R$/@',
                                           produto='Boi China', uf=_uf(praca), praca=praca))
    _exigir(len(registros) >= 10, 'Boi China: poucas linhas')
    return registros


def pagina_boi_gordo(doc) -> list[dict]:
    return _mercado_fisico(doc, 'Boi gordo', 'BOI GORDO') + _boi_china(doc)


def pagina_vaca_gorda(doc) -> list[dict]:
    return _mercado_fisico(doc, 'Vaca gorda', 'VACA GORDA')


def pagina_novilha(doc) -> list[dict]:
    return _mercado_fisico(doc, 'Novilha gorda', 'NOVILHA GORDA')


def pagina_indicadores(doc) -> list[dict]:
    tabela, m = _tabela(doc, r'^Indicador do boi gordo da Scot Consultoria \(R\$/@\) - (\d{2}/\d{2}/\d{4})$')
    data = _data(m.group(1))
    cabecalho = [_texto(c) for c in _celulas(_linhas(tabela)[1])]
    _exigir(cabecalho[1:3] == ['Hoje', 'Ontem'], 'colunas do indicador mudaram')
    registros = []
    for tr in _linhas(tabela, 'conteudo'):
        celulas = [_texto(c) for c in _celulas(tr)]
        original = celulas[0]
        praca = _limpar(original.replace('(kg)', ''))
        unidade = 'R$/kg vivo' if '(kg)' in original else 'R$/@'
        atual, anterior, variacao = _num(celulas[1]), _num(celulas[2]), _num(celulas[3])
        if atual is None:
            continue
        tendencia = _tendencia(tr)
        registros.append(_registro(data, 'Indicador do boi gordo', 'Indicador', atual, unidade,
                                   produto='Boi gordo', uf=_uf(praca), praca=praca, tendencia=tendencia))
        if anterior is not None:
            # A Scot mostra a variação sem sinal (o sinal está na seta); calcula hoje - ontem.
            diferenca = round(atual - anterior, 2)
            _exigir(variacao is None or abs(abs(diferenca) - variacao) < 0.015,
                    f'{praca}: variação {variacao} não confere com hoje - ontem')
            registros.append(_registro(data, 'Indicador do boi gordo', 'Variação', diferenca, unidade,
                                       produto='Boi gordo', uf=_uf(praca), praca=praca,
                                       tendencia=tendencia))
    _exigir(len({r['praca'] for r in registros}) >= 25, 'Indicador: poucas praças')
    return registros


def pagina_reposicao(doc) -> list[dict]:
    regex = re.compile(r'^((?:MACHO|F[ÊE]MEA)\b.*?) - (\d{2}/\d{2}/\d{4})$', re.IGNORECASE)
    registros, categorias = [], []
    for tabela in doc.iter('table'):
        celulas = tabela.xpath('.//th|.//td')
        m = regex.search(_texto(celulas[0])) if celulas else None
        if not m:
            continue
        categoria, data = m.group(1).capitalize(), _data(m.group(2))
        categorias.append(categoria)
        topo, titulo = _linhas(tabela, 'top'), _linhas(tabela, 'titulo')
        _exigir(len(topo) == 1 and len(titulo) == 1, f'{categoria}: cabeçalho da tabela mudou')
        animais = []
        for c in _celulas(topo[0])[1:]:
            nome = _texto(c)
            base, _, resto = nome.partition('(')
            animais.append(f'{base.strip().capitalize()} ({resto}' if resto else base.strip().capitalize())
        medidas = [_texto(c) for c in _celulas(titulo[0])]
        _exigir(medidas == ['UF'] + ['R$/cab', 'R$/kg', 'Troca'] * len(animais),
                f'{categoria}: colunas mudaram ({medidas})')
        for tr in _linhas(tabela, 'conteudo'):
            valores = [_texto(c) for c in _celulas(tr)]
            _exigir(len(valores) == 1 + 3 * len(animais), f'{categoria} {valores[:1]}: colunas a mais/menos')
            uf = _limpar(valores[0])
            praca = f'{uf} (cruzados)' if '*' in valores[0] else uf
            for i, animal in enumerate(animais):
                trio = valores[1 + 3 * i: 4 + 3 * i]
                for indicador, unidade, texto in zip(['R$/cab', 'R$/kg', 'Relação de troca'],
                                                     ['R$/cab', 'R$/kg', '@ de boi gordo'], trio):
                    valor = _num(texto)
                    if valor is not None:
                        registros.append(_registro(data, 'Reposição', indicador, valor, unidade,
                                                   categoria=categoria, produto=animal, uf=uf, praca=praca))
    _exigir({'Macho nelore', 'Fêmea nelore'} <= set(categorias),
            f'reposição sem as tabelas de nelore (achadas: {categorias})')
    return registros


def pagina_mercado_futuro(doc) -> list[dict]:
    tabela, m = _tabela(doc, r'^MERCADO FUTURO DO BOI GORDO - (\d{2}/\d{2}/\d{4})$')
    publicado = dt.date.fromisoformat(_data(m.group(1)))
    titulo = _linhas(tabela, 'titulo')
    _exigir(len(titulo) == 1, 'cabeçalho do mercado futuro mudou')
    cabecalho = [_texto(c) for c in _celulas(titulo[0])]
    _exigir(cabecalho[0] == 'Futuros' and cabecalho[3:] == ['C. Abertos', 'Variação', 'Câmbio', 'US$ à Vista'],
            f'colunas do mercado futuro mudaram ({cabecalho})')
    dia, mes = cabecalho[2].split('/')
    mes = MESES_ABREV[mes.lower()[:3]]
    ano = publicado.year - (1 if mes > publicado.month else 0)
    data = dt.date(ano, mes, int(dia)).isoformat()
    # A coluna "C. Abertos" da Scot traz, na verdade, os contratos negociados no dia
    # (confere com o boletim da B3: 2.621 no BGIV26 em 21/10/2026, não os 20.238 em aberto).
    campos = [('Ajuste anterior', 'R$/@'), ('Ajuste', 'R$/@'), ('Contratos negociados', 'contratos'),
              ('Variação', 'R$/@'), ('Câmbio', 'R$/US$'), ('US$ à vista', 'US$/@')]
    registros = []
    for tr in _linhas(tabela, 'conteudo'):
        celulas = [_texto(c) for c in _celulas(tr)]
        _exigir(re.match(r'^\w{3}/\d{2}$', celulas[0]) is not None, f'vencimento estranho: {celulas[0]}')
        for (indicador, unidade), texto in zip(campos, celulas[1:]):
            valor = _num(texto)
            if valor is not None:
                registros.append(_registro(data, 'Mercado futuro do boi', indicador, valor, unidade,
                                           categoria=celulas[0], produto='Boi gordo (B3)'))
    _exigir(len(registros) >= 12, 'mercado futuro: poucas linhas')
    return registros


def pagina_graos(doc) -> list[dict]:
    registros = []
    for produto in ('Milho', 'Soja'):
        tabela, m = _tabela(doc, rf'^{produto} - (\d{{2}}/\d{{2}}/\d{{4}})$')
        data = _data(m.group(1))
        cabecalho = [[_texto(c) for c in _celulas(tr)] for tr in _linhas(tabela)]
        _exigir(['UF', 'Cidade', 'Compra'] in cabecalho, f'colunas de {produto} mudaram')
        uf = ''
        for tr in _linhas(tabela, 'conteudo'):
            celulas = [_texto(c) for c in _celulas(tr)]
            uf = celulas[0] or uf
            valor = _num(celulas[2])
            if valor is not None:
                registros.append(_registro(data, produto, 'Compra', valor, 'R$/saca 60 kg',
                                           produto=produto, uf=uf, praca=_limpar(celulas[1])))
        _exigir(sum(r['grupo'] == produto for r in registros) >= 20, f'{produto}: poucas praças')
    return registros


def pagina_atacado(doc) -> list[dict]:
    tabela, m = _tabela(doc, r'^ATACADO - (\w{2}) - (\d{2}/\d{2}/\d{4})$')
    uf, data = m.group(1), _data(m.group(2))
    cabecalho = [_texto(c) for c in _celulas(_linhas(tabela)[0])]
    _exigir('R$/KG' in cabecalho[1].upper(), 'colunas do atacado mudaram')
    registros = []
    for tr in _linhas(tabela, 'conteudo'):
        celulas = [_texto(c) for c in _celulas(tr)]
        valor = _num(celulas[1])
        if valor is not None:
            registros.append(_registro(data, 'Atacado', 'Preço', valor, 'R$/kg', produto=_limpar(celulas[0]),
                                       uf=uf, praca='São Paulo' if uf == 'SP' else uf,
                                       tendencia=_tendencia(tr)))
    _exigir(len(registros) >= 10, 'atacado: poucas linhas')
    return registros


def pagina_equivalentes(doc) -> list[dict]:
    tabela, m = _tabela(doc, r'^EQUIVALENTES - (\d{2}/\d{2}/\d{4})$')
    data = _data(m.group(1))
    cabecalho = [_texto(c) for c in _celulas(_linhas(tabela)[1])]
    _exigir(cabecalho[1:] == ['BOI', 'MARGEM', 'VACA', 'MARGEM'], f'colunas dos equivalentes mudaram ({cabecalho})')
    registros = []
    for tr in _linhas(tabela, 'conteudo'):
        celulas = [_texto(c) for c in _celulas(tr)]
        item = _limpar(celulas[0])
        for i, produto in ((1, 'Boi'), (3, 'Vaca')):
            for indicador, unidade, texto in (('Valor', 'R$/@', celulas[i]), ('Margem', '%', celulas[i + 1])):
                valor = _num(texto)
                if valor is not None:
                    registros.append(_registro(data, 'Equivalentes', indicador, valor, unidade,
                                               categoria=item, produto=produto))
    _exigir(len(registros) >= 8, 'equivalentes: poucas linhas')
    return registros


def pagina_boi_no_mundo(doc) -> list[dict]:
    tabela, m = _tabela(doc, r'^BOI NO MUNDO - (\d{2}/\d{2}/\d{4})$')
    data = _data(m.group(1))
    cabecalho = [_texto(c) for c in _celulas(_linhas(tabela)[1])]
    _exigir(cabecalho[1].replace(' ', '') == 'US$/@', 'colunas do boi no mundo mudaram')
    registros = []
    for tr in _linhas(tabela, 'conteudo'):
        celulas = [_texto(c) for c in _celulas(tr)]
        valor = _num(celulas[1])
        if valor is not None:
            registros.append(_registro(data, 'Boi no mundo', 'Preço', valor, 'US$/@', produto='Boi gordo',
                                       praca=_limpar(celulas[0])))
    _exigir(len(registros) >= 5, 'boi no mundo: poucas linhas')
    return registros


def pagina_couro_e_sebo(doc) -> list[dict]:
    tabela, m = _tabela(doc, r'^COURO E SEBO - (\d{2}/\d{2}/\d{4})$')
    data = _data(m.group(1))
    cabecalho = [_texto(c) for c in _celulas(_linhas(tabela)[1])]
    _exigir('R$/KG' in cabecalho[0].upper() and len(cabecalho) == 3, 'colunas do couro e sebo mudaram')
    pracas = [p.title().replace(' Do ', ' do ') for p in cabecalho[1:]]
    registros = []
    for tr in _linhas(tabela, 'conteudo'):
        celulas = [_texto(c) for c in _celulas(tr)]
        for praca, texto in zip(pracas, celulas[1:]):
            valor = _num(texto)
            if valor is not None:
                registros.append(_registro(data, 'Couro e sebo', 'Preço', valor, 'R$/kg',
                                           produto=_limpar(celulas[0]), uf=_uf(praca), praca=praca))
    _exigir(len(registros) >= 3, 'couro e sebo: poucas linhas')
    return registros


def pagina_leite(doc) -> list[dict]:
    regex = re.compile(r'pago aos produtores em (\w+)/(\d{2}), pela produ\w+ de (\w+)/(\d{2})', re.IGNORECASE)
    m = next((regex.search(_texto(el)) for el in doc.iter('td', 'th', 'p', 'div', 'span', 'strong')
              if regex.search(_texto(el))), None)
    _exigir(m is not None, 'título do leite ("pago aos produtores em ...") não encontrado')
    data = dt.date(2000 + int(m.group(2)), MESES[m.group(1).lower()], 1).isoformat()
    campos = [('Qualidade mínimo', 'R$/L'), ('Qualidade máximo', 'R$/L'), ('Padrão mínimo', 'R$/L'),
              ('Padrão máximo', 'R$/L'), ('Média qualidade', 'R$/L'), ('Média qualidade', 'US$/L'),
              ('Média padrão', 'R$/L'), ('Média padrão', 'US$/L')]
    registros, regioes = [], 0
    # Cada região é uma tabela de uma linha (região + 8 valores) dentro da linha da UF.
    for tabela in doc.iter('table'):
        trs = _linhas(tabela)
        if len(trs) != 1:
            continue
        celulas = [_texto(c) for c in _celulas(trs[0])]
        if len(celulas) != 9 or not re.match(r'^\d+,\d{3}$', celulas[1]):
            continue
        linha_uf = tabela.getparent().getparent()
        uf = _texto(_celulas(linha_uf)[0]) if linha_uf is not None and linha_uf.tag == 'tr' else ''
        _exigir(uf in UFS, f'leite: UF não encontrada para {celulas[0]}')
        regioes += 1
        for (indicador, unidade), texto in zip(campos, celulas[1:]):
            valor = _num(texto)
            if valor:  # 0,000 = sem levantamento
                registros.append(_registro(data, 'Leite', indicador, valor, unidade, produto='Leite',
                                           uf=uf, praca=_limpar(celulas[0])))
    _exigir(regioes >= 20, f'leite: só {regioes} regiões')
    return registros


PAGINAS = {
    'boi-gordo': pagina_boi_gordo,
    'vaca-gorda': pagina_vaca_gorda,
    'novilha': pagina_novilha,
    'indicadores': pagina_indicadores,
    'reposicao': pagina_reposicao,
    'mercado-futuro': pagina_mercado_futuro,
    'graos': pagina_graos,
    'atacado': pagina_atacado,
    'equivalentes': pagina_equivalentes,
    'boi-no-mundo': pagina_boi_no_mundo,
    'couro-e-sebo': pagina_couro_e_sebo,
    'leite-cotacoes': pagina_leite,
}


# --------------------------------------------------------------------------- persistência

def _gravar_json(destino: Path, dados: dict) -> None:
    destino.parent.mkdir(parents=True, exist_ok=True)
    tmp = destino.with_suffix('.tmp')
    tmp.write_text(json.dumps(dados, indent=2, sort_keys=True, ensure_ascii=False), encoding='utf-8')
    tmp.replace(destino)


def carregar_controle() -> dict:
    try:
        return json.loads(ARQ_CONTROLE.read_text(encoding='utf-8'))
    except (FileNotFoundError, json.JSONDecodeError):
        return {}


def _ler_csv(arquivo: Path) -> list[dict]:
    with open(arquivo, newline='', encoding='utf-8') as f:
        return [{**l, 'origem': l.get('origem') or ORIGEM_PROPRIA} for l in csv.DictReader(f)]


def _ordem_grupo(grupo: str) -> int:
    return GRUPOS.index(grupo) if grupo in GRUPOS else len(GRUPOS)


def _texto_valor(valor: float) -> str:
    return repr(round(valor, 6))


def gravar(registros: list[dict], controle: dict, coletado_em: str) -> list[str]:
    """Grava os grupos novos ou alterados nos CSVs diários; devolve o que mudou.
    Registros com 'coletado_em' e 'origem' próprios (histórico importado) os mantêm."""
    mudancas = []
    por_data = collections.defaultdict(list)
    for r in registros:
        por_data[r['data']].append({**r, 'valor': _texto_valor(r['valor']),
                                    'coletado_em': r.get('coletado_em', coletado_em),
                                    'origem': r.get('origem', ORIGEM_PROPRIA)})
    sem_coleta = lambda linhas: [{k: v for k, v in l.items() if k != 'coletado_em'} for l in linhas]

    for data, novos in sorted(por_data.items()):
        arquivo = DIR_DIARIO / f'{data}.csv'
        linhas = _ler_csv(arquivo) if arquivo.exists() else []
        alterou = False
        for grupo in sorted({r['grupo'] for r in novos}, key=_ordem_grupo):
            do_grupo = [r for r in novos if r['grupo'] == grupo]
            antigos = [l for l in linhas if l['grupo'] == grupo]
            if sem_coleta(antigos) == sem_coleta(do_grupo):
                continue
            linhas = [l for l in linhas if l['grupo'] != grupo] + do_grupo
            alterou = True
            reg = controle.setdefault(data, {}).get(grupo)
            if reg:
                reg.update(linhas=len(do_grupo), ultima_alteracao=coletado_em, versoes=reg.get('versoes', 1) + 1)
                mudancas.append(f'{grupo} {data}: revisado pela Scot ({len(do_grupo)} linhas, versão {reg["versoes"]})')
            else:
                controle[data][grupo] = {'linhas': len(do_grupo), 'primeira_coleta': do_grupo[0]['coletado_em'],
                                         'ultima_alteracao': coletado_em, 'versoes': 1}
                if do_grupo[0]['origem'] != ORIGEM_PROPRIA:
                    controle[data][grupo]['origem'] = do_grupo[0]['origem']
                mudancas.append(f'{grupo} {data}: novo ({len(do_grupo)} linhas)')
        if alterou:
            linhas.sort(key=lambda l: _ordem_grupo(l['grupo']))
            DIR_DIARIO.mkdir(parents=True, exist_ok=True)
            tmp = arquivo.with_suffix('.tmp')
            with open(tmp, 'w', newline='', encoding='utf-8') as f:
                w = csv.DictWriter(f, fieldnames=CAMPOS_CSV)
                w.writeheader()
                w.writerows(linhas)
            tmp.replace(arquivo)
    if mudancas:
        _gravar_json(ARQ_CONTROLE, controle)
    return mudancas


def ler_todos_os_dias() -> list[dict]:
    linhas = []
    for arquivo in sorted(DIR_DIARIO.glob('*.csv')):
        linhas.extend(_ler_csv(arquivo))
    # Ordenação estável: dentro da data e do grupo, fica a ordem do site.
    return sorted(linhas, key=lambda l: (l['data'], _ordem_grupo(l['grupo'])))


# --------------------------------------------------------------------------- saídas

def arquivo_csv(ano: int) -> Path:
    return ARQ_CSV.with_name(f'{ARQ_CSV.stem}_{ano}.csv')


def gerar_csv(linhas: list[dict]) -> str:
    """Formato longo (uma linha por data, praça e indicador), vírgula, ponto decimal,
    datas AAAA-MM-DD, para o Power BI ou outro programa ler direto do GitHub.
    Um arquivo por ano (~50 MB/ano; o GitHub recusa arquivo acima de 100 MB), do
    primeiro ano até o atual, mesmo que ainda vazio."""
    ARQ_CSV.parent.mkdir(parents=True, exist_ok=True)
    por_ano = collections.defaultdict(list)
    for l in linhas:
        por_ano[int(l['data'][:4])].append(l)
    anos = range(min(por_ano, default=hoje().year), hoje().year + 1)
    for ano in anos:
        destino = arquivo_csv(ano)
        tmp = destino.with_suffix('.tmp')
        with open(tmp, 'w', newline='', encoding='utf-8') as f:
            w = csv.writer(f)
            w.writerow(COLUNAS_CSV)
            for l in por_ano.get(ano, []):
                w.writerow([l[c] for c in CAMPOS_CSV])
        tmp.replace(destino)
    return f'{arquivo_csv(anos[0]).parent / (ARQ_CSV.stem + "_AAAA.csv")} ({anos[0]} a {anos[-1]})'


def _nome_tabela(grupo: str) -> str:
    sem_acento = unicodedata.normalize('NFKD', grupo).encode('ascii', 'ignore').decode()
    return ''.join(p.capitalize() for p in re.split(r'\W+', sem_acento) if p)


def _aba_grupo(wb: Workbook, grupo: str, linhas: list[dict]) -> None:
    """Uma aba por grupo, com uma coluna por indicador (formato largo)."""
    chaves = ['data', 'categoria', 'produto', 'uf', 'praca']
    chaves = [c for c in chaves if c == 'data' or len({l[c] for l in linhas}) > 1 or
              (c != 'produto' and any(l[c] for l in linhas))]
    nomes = {'data': 'Data', 'categoria': 'Categoria', 'produto': 'Produto', 'uf': 'UF', 'praca': 'Praça'}
    indicadores = list(dict.fromkeys(l['indicador'] for l in linhas))
    unidades = {i: list(dict.fromkeys(l['unidade'] for l in linhas if l['indicador'] == i)) for i in indicadores}
    por_linha = collections.defaultdict(set)
    for l in linhas:
        por_linha[(l['indicador'],) + tuple(l[c] for c in chaves)].add(l['unidade'])
    # Indicador com uma unidade só leva a unidade no cabeçalho. Se a unidade muda de uma
    # linha para outra (ex.: RS em R$/kg vivo), entra a coluna Unidade. Se a mesma linha tem
    # duas unidades (Leite: média em R$/L e em US$/L), cada unidade vira uma coluna.
    varias_na_linha = {k[0] for k, u in por_linha.items() if len(u) > 1}
    variam = [i for i in indicadores if len(unidades[i]) > 1 and i not in varias_na_linha]

    colunas = []
    for i in indicadores:
        for u in (unidades[i] if i not in variam else [None]):
            colunas.append((i, u))
    linhas_largas: dict[tuple, dict] = {}
    for l in linhas:
        chave = tuple(l[c] for c in chaves)
        item = linhas_largas.setdefault(chave, {'unidade': '', 'tendencia': '', 'origem': l['origem']})
        coluna = (l['indicador'], None if l['indicador'] in variam else l['unidade'])
        item[coluna] = float(l['valor'])
        if l['indicador'] in variam and not item['unidade']:
            item['unidade'] = l['unidade']
        item['tendencia'] = item['tendencia'] or l['tendencia']
    tem_tendencia = any(l['tendencia'] for l in linhas)
    tem_origem = len({l['origem'] for l in linhas}) > 1

    cabecalho = [nomes[c] for c in chaves] + (['Unidade'] if variam else [])
    cabecalho += [i if u in (None, i) else f'{i} ({u})' for i, u in colunas]
    cabecalho += ['Tendência'] if tem_tendencia else []
    cabecalho += ['Origem'] if tem_origem else []
    ws = wb.create_sheet(grupo[:31])
    ws.append(cabecalho)
    for chave, item in linhas_largas.items():
        valores = [dt.datetime.strptime(v, '%Y-%m-%d') if c == 'data' else v for c, v in zip(chaves, chave)]
        valores += [item['unidade']] if variam else []
        valores += [item.get(c) for c in colunas]
        valores += [item['tendencia']] if tem_tendencia else []
        valores += [item['origem']] if tem_origem else []
        ws.append(valores)

    n = max(len(linhas_largas) + 1, 2)
    inicio_valores = len(chaves) + (2 if variam else 1)
    for j, cab in enumerate(cabecalho, 1):
        letra = get_column_letter(j)
        formato = ('dd/mm/yyyy' if j == 1 else
                   '#,##0.00##' if inicio_valores <= j < inicio_valores + len(colunas) else '@')
        for (cel,) in ws[f'{letra}2:{letra}{n}']:
            cel.number_format = formato
        ws.column_dimensions[letra].width = max(min(len(cab) + 3, 30), 11)
    for c in ws[1]:
        c.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 30
    tabela = Table(displayName=_nome_tabela(grupo), ref=f'A1:{get_column_letter(len(cabecalho))}{n}')
    tabela.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tabela)
    ws.freeze_panes = ws.cell(row=2, column=inicio_valores)


def gerar_excel(linhas: list[dict]) -> Path:
    wb = Workbook()
    wb._named_styles['Normal'].font = Font(name='Arial', size=10)
    info = wb.active
    info.title = 'Sobre'
    grupos = sorted({l['grupo'] for l in linhas}, key=_ordem_grupo)
    for grupo in grupos:
        _aba_grupo(wb, grupo, [l for l in linhas if l['grupo'] == grupo])

    info.column_dimensions['A'].width = 24
    info.column_dimensions['B'].width = 120
    negrito = Font(name='Arial', size=10, bold=True)
    datas = sorted({l['data'] for l in linhas if l['grupo'] == 'Boi gordo'})
    meta = [
        ('Período', f'{dt.date.fromisoformat(datas[0]):%d/%m/%Y} a {dt.date.fromisoformat(datas[-1]):%d/%m/%Y} '
                    f'({len(datas)} data{"s" if len(datas) > 1 else ""} do boi gordo)' if datas else 'sem dados'),
        ('Atualizado em', f'{agora():%d/%m/%Y %H:%M} (horário de Brasília)'),
        ('Como atualizar', 'Atualizada pelo GitHub Actions em github.com/Murkiriel/scot-cotacoes (veja Releases). '
                           'Localmente: python scot_cotacoes.py'),
        ('Data', 'A data é a da cotação publicada pela Scot, não a da coleta (no Leite, o mês do pagamento)'),
        ('Histórico', 'O site só mostra a cotação mais recente. Coleta própria desde 22/10/2026; antes, '
                      'histórico importado de outros projetos (boi gordo e novilha desde 26/12/2025, reposição '
                      'de machos nelore desde 23/02/2026), com a data corrigida: coluna Origem'),
        ('', ''),
        ('Grupo', 'Notas da Scot'),
    ]
    meta += [(g, NOTAS.get(g, '')) for g in grupos]
    meta += [('', ''), ('Coluna do CSV', 'Descrição')]
    meta += list(zip(COLUNAS_CSV, DESCRICAO_COLUNAS))
    for a, b in meta:
        info.append([a, b])
    for linha in info.iter_rows():
        linha[0].font = negrito
        linha[1].alignment = Alignment(wrap_text=True, vertical='top')
    wb.move_sheet(info, offset=-len(wb.sheetnames))

    try:
        wb.save(ARQ_EXCEL)
        return ARQ_EXCEL
    except PermissionError:
        alternativo = ARQ_EXCEL.with_name(f'{ARQ_EXCEL.stem}_{dt.datetime.now():%Y%m%d_%H%M%S}.xlsx')
        wb.save(alternativo)
        log(f'AVISO: {ARQ_EXCEL.name} está aberto no Excel; planilha salva como {alternativo.name}')
        return alternativo


# --------------------------------------------------------------------------- execução

def grupos_atrasados(linhas: list[dict]) -> list[str]:
    ultima = {}
    for l in linhas:
        ultima[l['grupo']] = max(ultima.get(l['grupo'], ''), l['data'])
    atrasados = []
    for grupo, data in sorted(ultima.items(), key=lambda g: _ordem_grupo(g[0])):
        dias = (hoje() - dt.date.fromisoformat(data)).days
        if dias > ATRASO_MAXIMO_DIAS_GRUPO.get(grupo, ATRASO_MAXIMO_DIAS):
            atrasados.append(f'{grupo} (última data {dt.date.fromisoformat(data):%d/%m/%Y}, {dias} dias)')
    return atrasados


def coletar(pagina: str, pasta_html: Path | None) -> list[dict]:
    html = ((pasta_html / f'{pagina}.html').read_text(encoding='utf-8') if pasta_html
            else baixar(pagina))
    doc = lxml.html.fromstring(html)
    try:
        registros = PAGINAS[pagina](doc)
    except ValueError as e:
        # O site às vezes devolve outra página (vista em 23/10/2026 no couro e sebo, que
        # minutos antes e depois veio normal): registra o que chegou para diagnóstico.
        titulo = ' '.join(''.join(doc.xpath('//title//text()')).split())[:80]
        raise ValueError(f'{e} [recebido: {len(html):,} caracteres, título "{titulo}"]') from e
    futuras = sorted({r['data'] for r in registros if r['data'] > (hoje() + dt.timedelta(days=1)).isoformat()})
    _exigir(not futuras or pasta_html is not None, f'data no futuro: {futuras}')
    return registros


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('--paginas', nargs='+', default=list(PAGINAS), choices=list(PAGINAS),
                    metavar='PAGINA', help='páginas a coletar, padrão todas: ' + ' '.join(PAGINAS))
    ap.add_argument('--html', type=Path, help='lê <PAGINA>.html desta pasta em vez de baixar do site')
    ap.add_argument('--sem-gravar', action='store_true',
                    help='só lê e confere as páginas, sem gravar nada (para testar)')
    ap.add_argument('--sem-excel', action='store_true',
                    help='não gera a planilha .xlsx (só o CSV consolidado)')
    args = ap.parse_args()

    controle = carregar_controle()
    coletado_em = agora().isoformat()
    erros, mudancas = [], []
    for n, pagina in enumerate(args.paginas):
        if n and not args.html:
            time.sleep(PAUSA_ENTRE_PAGINAS_S)
        try:
            registros = coletar(pagina, args.html)
        except Bloqueado as e:
            erros += args.paginas[n:]
            log(f'{pagina}: BLOQUEADO pelo site ({e}); as {len(args.paginas) - n} páginas restantes '
                'ficam para a próxima coleta')
            break
        except Exception as e:  # uma página com problema não impede as demais
            erros.append(pagina)
            log(f'{pagina}: ERRO {type(e).__name__}: {e}')
            continue
        resumo = collections.Counter((r['grupo'], r['data']) for r in registros)
        log(f'{pagina}: ' + '; '.join(f'{g} {dt.date.fromisoformat(d):%d/%m/%Y} ({q} linhas)'
                                      for (g, d), q in resumo.items()))
        if not args.sem_gravar:
            mudancas += gravar(registros, controle, coletado_em)

    if args.sem_gravar:
        log(f'Teste concluído; {len(erros)} página(s) com erro')
        return 1 if erros else 0
    for m in mudancas:
        log(f'  {m}')
    log(f'{len(mudancas)} grupo(s) novo(s) ou revisado(s); {len(erros)} página(s) com erro')
    linhas = ler_todos_os_dias()
    log(f'CSV: {gerar_csv(linhas)} ({len(linhas)} linhas)')
    if not args.sem_excel:
        log(f'Planilha: {gerar_excel(linhas)}')
    if erros:
        log('Páginas com erro (serão tentadas de novo na próxima execução): ' + ', '.join(erros))
    atrasados = [] if args.html else grupos_atrasados(linhas)
    if atrasados:
        log('ATRASO: grupos sem data nova há muito tempo: ' + '; '.join(atrasados))
        return 2
    return 1 if erros else 0


if __name__ == '__main__':
    sys.exit(main())
